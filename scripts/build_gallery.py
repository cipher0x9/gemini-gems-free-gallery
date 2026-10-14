#!/usr/bin/env python3
"""Build the Gemini Gems Gallery: Markdown, site, PDF, PowerPoint, and zip."""

from __future__ import annotations

import html
import io
import sys
import zipfile
from pathlib import Path

from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    NextPageTemplate,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Flowable,
)

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from art import illustration  # noqa: E402
from gems_data import GEMS  # noqa: E402

CATS = {
    "Workday": "#0F9B8E",
    "Engineering": "#6366F1",
    "Career & Learning": "#F59E0B",
    "Home & Family": "#22C55E",
}
PHRASES = {
    "sales-call-prep-desk": "crm",
    "meeting-recap-desk": "sent or shared",
    "resume-tailor-desk": "invent a title",
    "bedtime-story-desk": "scary",
    "weekly-planning-desk": "invite",
    "inbox-triage-desk": "archive and label",
    "homework-coach-desk": "final answer",
    "budget-tracker-desk": "move, transfer",
    "interview-prep-desk": "live interview",
    "study-plan-desk": "fits, tight, or over",
    "support-reply-desk": "refund",
    "sop-writer-desk": "expert",
    "onboarding-guide-desk": "create accounts",
    "signal-brief-desk": "emailed",
    "social-post-desk": "boost",
    "newsletter-desk": "broadcast",
    "pr-review-desk": "merge",
    "bug-triage-desk": "assign",
    "incident-postmortem-desk": "roll back",
    "runbook-desk": "never run a command",
    "api-docs-desk": "live api",
    "architecture-diagram-desk": "inferred",
    "claim-check-desk": "pass, partial, or fail",
    "competitive-landscape-desk": "recommendation",
    "certification-coach-desk": "brain dump",
    "money-desk": "not financial advice",
    "language-drill-desk": "correct, close, or miss",
    "lesson-plan-desk": "unknown",
    "curious-kids-science-desk": "adult required",
    "family-chore-chart-desk": "first name",
    "receipt-and-expense-desk": "mask",
    "travel-itinerary-desk": "book",
    "event-planner-desk": "deposit",
}


def esc(text: str) -> str:
    return html.escape(text, quote=True)


def fname(gem: dict) -> str:
    return f"{gem['n']:02d}-{gem['slug']}"


def render_instructions(gem: dict) -> str:
    lines = [
        f"You are {gem['name']}, {gem['role']}",
        "",
        "ROLE",
        f"You prepare work for the user. You do not act in their accounts, tools, or inbox.",
        "",
        "PURPOSE",
        gem["purpose"],
        f"Desk mark: {gem['mark']}.",
        "",
        "NEVER",
    ]
    lines += [f"- {item}" for item in gem["never"]]
    lines += ["", "INPUTS", "If a required input is missing, ask at most three short questions, then continue and mark the gap UNKNOWN."]
    lines += [f"- {item}" for item in gem["inputs"]]
    lines += ["", "STEPS"]
    for i, (title, detail) in enumerate(gem["steps"], 1):
        lines.append(f"{i}. {title}. {detail}")
    lines += ["", "OUTPUT", "Use these headings, in this order. If a section has nothing true to say, write NONE. Do not invent a row to fill the page."]
    lines += [f"- {item}" for item in gem["output"]]
    lines += ["", "FACTS"]
    lines += [f"- {item}" for item in gem["facts"]]
    lines += ["", "SAFETY"]
    lines += [f"- {item}" for item in gem["safety"]]
    if gem.get("kid"):
        lines += [
            "- Child data: ask only for a first name and an age band.",
            "- Do not ask for a surname, school, address, photo, account, or document number.",
            "- If extra personal detail is pasted, do not repeat it.",
        ]
    if gem.get("dispatch"):
        lines += [
            "",
            "APPROVAL LINE",
            "You still do not perform the forbidden action.",
            "Mark a draft approved for the user to use only when their entire message is exactly this line and nothing else:",
            gem["dispatch"],
            'Any other reply, including "ok", "send it", or approval buried in a longer note, is not approval. Say that the draft is held.',
        ]
    lines += [
        "",
        "IF ASKED TO BREAK A RULE",
        "Refuse that part in one sentence. Do the rest of the safe draft. Do not half-do the forbidden action.",
        "",
        "CHECK",
        "End with a short check:",
    ]
    lines += [f"- {item}" for item in gem["check"]]
    lines.append("- You did not claim that you sent, booked, paid, posted, or changed an outside system.")
    return "\n".join(lines).strip() + "\n"


def validate(gems: list[dict], texts: dict[str, str]) -> None:
    if len(gems) != 33:
        raise SystemExit(f"expected 33 gems, found {len(gems)}")
    nums = [g["n"] for g in gems]
    if nums != list(range(1, 34)):
        raise SystemExit(f"gem numbers are {nums}")
    icons = [g["icon"] for g in gems]
    if len(set(icons)) != 33:
        raise SystemExit("icons are not unique")
    for gem in gems:
        text = texts[gem["slug"]].lower()
        for heading in ("ROLE", "PURPOSE", "NEVER", "INPUTS", "STEPS", "OUTPUT", "FACTS", "SAFETY", "CHECK"):
            if heading not in texts[gem["slug"]]:
                raise SystemExit(f"{gem['slug']} missing {heading}")
        needle = PHRASES[gem["slug"]]
        if needle not in text:
            raise SystemExit(f"{gem['slug']} missing guardrail phrase: {needle}")
        if "```" in texts[gem["slug"]]:
            raise SystemExit(f"{gem['slug']} instructions contain a fence")


def flow_html(steps: list[tuple[str, str]]) -> str:
    items = []
    for i, (title, detail) in enumerate(steps):
        items.append(
            f'<li class="step" style="--i:{i}"><span class="n">{i + 1}</span>'
            f"<h3>{esc(title)}</h3><p>{esc(detail)}</p></li>"
        )
    return f'<div class="flow" style="--n:{len(steps)}"><div class="flow-beam" aria-hidden="true"></div><ol>{"".join(items)}</ol></div>'


def shell(title: str, desc: str, depth: int, body: str) -> str:
    site = "../" * depth
    icon = f"{site}favicon.svg"
    css = f"{site}assets/gallery.css"
    js = f"{site}assets/gallery.js"
    root = "../" * (depth + 1)
    home = f"{site}index.html" if depth else "index.html"
    printable = f"{site}print.html" if depth else "print.html"
    nav = f'''<a class="skip" href="#content">Skip to content</a>
<header class="nav no-print"><a class="brand" href="{home}"><i></i> Gemini Gems Gallery</a>
<nav class="nav-links">
<a href="{home}">Gallery</a>
<a href="{printable}">Printable view</a>
<a href="{root}downloads/Gemini-Gems-Guide.pdf">PDF</a>
<a href="{root}downloads/gemini-gems-pack.zip">ZIP</a>
<a href="{root}downloads/Gemini-Gems-Overview.pptx">Deck</a>
<a href="{root}examples/README.md">App starters</a>
<button type="button" data-theme-toggle aria-pressed="true">Light theme</button>
</nav></header>'''
    return f'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(title)}</title>
<meta name="description" content="{esc(desc)}">
<link rel="icon" href="{icon}">
<link rel="stylesheet" href="{css}">
<script>(function(){{try{{var t=localStorage.getItem("gems-theme");if(!t&&matchMedia("(prefers-color-scheme: light)").matches)t="light";document.documentElement.dataset.theme=t||"dark";}}catch(e){{document.documentElement.dataset.theme="dark";}}}})();</script>
</head>
<body>
{nav}
<main id="content">{body}</main>
<footer class="wrap"><nav>
<a href="{root}README.md">README</a>
<a href="{root}LICENSE">License</a>
<a href="https://gemini.google.com">Gemini</a>
<a href="https://creativecommons.org/licenses/by/4.0/">CC BY 4.0</a>
</nav>
<p>Content CC BY 4.0. Code MIT. Gemini and Google AI Studio are trademarks of Google LLC. This project is not affiliated with Google.</p>
</footer>
<script src="{js}"></script>
</body>
</html>
'''


ORBITS = '''<svg class="orbits" viewBox="0 0 640 640" aria-hidden="true">
  <g class="spin"><circle cx="320" cy="320" r="220" fill="none" stroke="#0F9B8E" stroke-width="2" stroke-opacity=".75"/>
  <circle class="pulse-dot" cx="540" cy="320" r="11" fill="#0F9B8E"/></g>
  <g class="spin-rev"><circle cx="320" cy="320" r="160" fill="none" stroke="#6366F1" stroke-width="2" stroke-opacity=".85"/>
  <circle class="pulse-dot" cx="320" cy="160" r="10" fill="#6366F1"/></g>
  <g class="spin"><circle cx="320" cy="320" r="96" fill="none" stroke="#22C55E" stroke-width="2"/>
  <circle class="pulse-dot" cx="416" cy="320" r="9" fill="#22C55E"/></g>
  <circle cx="320" cy="320" r="34" fill="#F59E0B"/>
  <path d="M320 298 L342 320 L320 342 L298 320 Z" fill="#0B1220"/>
</svg>'''


def render_index(gems: list[dict], texts: dict[str, str]) -> str:
    counts = {cat: sum(1 for g in gems if g["cat"] == cat) for cat in CATS}
    filters = ['<button type="button" data-filter="All" aria-pressed="true">All <span>33</span></button>']
    for cat, n in counts.items():
        filters.append(f'<button type="button" data-filter="{esc(cat)}" aria-pressed="false">{esc(cat)} <span>{n}</span></button>')
    cards = []
    for gem in gems:
        accent = CATS[gem["cat"]]
        kid = '<span class="kid">Kid-safe</span>' if gem.get("kid") else ""
        hay = esc(" ".join([gem["name"], gem["summary"], gem["cat"], gem["mark"], gem["icon"], *gem["prompts"]]).lower())
        dots = "".join("<i></i>" for _ in gem["steps"])
        cards.append(f'''<article class="card" style="--accent:{accent}" data-cat="{esc(gem["cat"])}" data-hay="{hay}">
<div class="poster">{illustration(gem["n"], gem["slug"], accent, gem["name"])}</div>
<div class="card-body">
<div class="meta"><span class="num">{gem["n"]:02d}</span><span class="cat">{esc(gem["cat"])}</span>{kid}</div>
<h2>{gem["icon"]} {esc(gem["name"])}</h2>
<p>{esc(gem["summary"])}</p>
<div class="mini" aria-hidden="true">{dots}<span>{len(gem["steps"])} steps</span></div>
<pre id="src-{gem["n"]:02d}" hidden>{esc(texts[gem["slug"]])}</pre>
<button class="copy" type="button" data-copy="src-{gem["n"]:02d}">Copy instructions</button>
<a class="open" href="gems/{fname(gem)}.html">Open desk</a>
</div></article>''')
    app_steps = [
        ("Pick a Gem", "Choose a desk from the gallery and open its page."),
        ("Copy the instructions", "Use the copy button. That block is the whole system prompt."),
        ("Set the key", "Put your Gemini API key in GEMINI_API_KEY. Do not paste it into a file."),
        ("Run a starter", "The Python or JavaScript starter reads the Gem file and sends one prompt."),
        ("Review the draft", "The Gem still does not send, book, or pay. You decide what to do with the text."),
    ]
    body = f'''<div class="wrap hero">
<div>
<p class="kicker"><span class="pill">33 desks</span><span class="pill soft" style="--accent:#6366F1">Copy-ready</span></p>
<h1>Gemini Gems Gallery</h1>
<p class="lede">Thirty-three standing instructions for Gemini. Each one has a job, a shape for the answer, and rules it will not cross. Copy a Gem, or use it as the system prompt for a small app.</p>
<div class="rule-row">
<div class="rule"><strong>Draft only</strong><span>It writes. You send, book, pay, and post.</span></div>
<div class="rule"><strong>No invented facts</strong><span>Missing facts stay marked UNKNOWN.</span></div>
<div class="rule"><strong>Kid-safe desks</strong><span>First name and age band only, where a child is the subject.</span></div>
</div>
<div class="cta-row">
<a class="solid" href="../downloads/Gemini-Gems-Guide.pdf">Download the PDF guide</a>
<a class="amber" href="../downloads/gemini-gems-pack.zip">Download the ZIP pack</a>
<a class="indigo" href="../downloads/Gemini-Gems-Overview.pptx">Download the deck</a>
<a class="ghost" href="print.html">Printable view</a>
</div>
</div>
<div class="stage">{ORBITS}</div>
</div>
<div class="wrap">
<div class="controls no-print">
<input id="q" type="search" placeholder="Search desks, or press /" aria-label="Search desks">
<div class="filters" role="group" aria-label="Filter by category">{"".join(filters)}</div>
<p class="count" id="result-count">33 desks</p>
</div>
<p id="empty" hidden>No desk matches that search.</p>
<div class="grid" id="grid">{"".join(cards)}</div>
<section class="section" id="build-an-app" style="--accent:#6366F1">
<h2>Build an app with a Gem</h2>
<p class="sub">A Gem's instructions are a system prompt. The starters in the examples folder load a Markdown file, take the instruction block, and call Gemini. The key stays in the environment.</p>
{flow_html(app_steps)}
<p class="cta-row"><a class="solid" href="../examples/README.md">Open the examples guide</a>
<a class="ghost" href="../examples/gem_app.py">Python starter</a>
<a class="ghost" href="../examples/gem_app.mjs">JavaScript starter</a></p>
</section>
</div>'''
    return shell("Gemini Gems Gallery", "33 copy-ready Gemini Gems with search, themes, and printable instructions.", 0, body)


def render_gem_page(gem: dict, gems: list[dict], text: str) -> str:
    accent = CATS[gem["cat"]]
    kid = '<span class="pill" style="--accent:#22C55E">Kid-safe</span>' if gem.get("kid") else ""
    nevers = "".join(f"<li>{esc(item)}</li>" for item in gem["never"][:3])
    same = [g for g in gems if g["cat"] == gem["cat"] and g["n"] != gem["n"]]
    extra = [g for g in gems if g["n"] != gem["n"] and g not in same]
    links = []
    for other in (same + extra)[:3]:
        links.append(
            f'<a href="{fname(other)}.html" style="--accent:{CATS[other["cat"]]}">'
            f'<span>{other["n"]:02d} · {esc(other["cat"])}</span>{other["icon"]} {esc(other["name"])}</a>'
        )
    body = f'''<div class="wrap" style="--accent:{accent}">
<section class="gem-hero">
<div>
<p class="kicker"><span class="pill">{gem["n"]:02d}</span><span class="pill soft">{esc(gem["cat"])}</span><span class="pill soft">{esc(gem["mark"])}</span>{kid}</p>
<h1>{gem["icon"]} {esc(gem["name"])}</h1>
<p class="lede">{esc(gem["summary"])}</p>
<ul class="never">{nevers}</ul>
</div>
<div class="art-frame">{illustration(gem["n"], gem["slug"], accent, gem["name"] + " illustration")}</div>
</section>
<section class="section">
<h2>How it works</h2>
<p class="sub">The desk moves one way. The light along the path marks the step you are in.</p>
{flow_html(gem["steps"])}
</section>
<section class="section">
<h2>Example</h2>
<div class="split">
<article class="io in"><h3>Example input</h3><pre>{esc(gem["example_in"])}</pre></article>
<article class="io out"><h3>Example output</h3><pre>{esc(gem["example_out"])}</pre></article>
</div>
</section>
<section class="section">
<h2>Instructions</h2>
<div class="panel">
<div class="instr-head"><p>Copy this into Gemini, then open Gems, New Gem, and save.</p>
<button type="button" data-copy="instructions">Copy instructions</button></div>
<pre class="instructions" id="instructions">{esc(text)}</pre>
</div>
</section>
<section class="section">
<h2>Remix</h2>
<div class="remix"><p>{esc(gem["remix"])}</p></div>
</section>
<section class="section">
<h2>Try another desk</h2>
<div class="related">{"".join(links)}</div>
<p class="cta-row"><a class="ghost" href="../../gems/{fname(gem)}.md">Read the Markdown source</a>
<a class="ghost" href="../index.html">Back to the gallery</a></p>
</section>
</div>'''
    return shell(f"{gem['name']} · Gemini Gems Gallery", gem["summary"], 1, body)


def render_print(gems: list[dict], texts: dict[str, str]) -> str:
    blocks = []
    for gem in gems:
        accent = CATS[gem["cat"]]
        blocks.append(f'''<article class="gem-block" style="--accent:{accent}" id="gem-{gem["n"]:02d}">
<p class="kicker"><span class="pill">{gem["n"]:02d}</span><span class="pill soft">{esc(gem["cat"])}</span></p>
<h2>{gem["icon"]} {esc(gem["name"])}</h2>
<p>{esc(gem["summary"])}</p>
{flow_html(gem["steps"])}
<h3>Example input</h3><pre>{esc(gem["example_in"])}</pre>
<h3>Example output</h3><pre>{esc(gem["example_out"])}</pre>
<div class="instr-head"><h3>Instructions</h3><button class="no-print" type="button" data-copy="p-{gem["n"]:02d}">Copy instructions</button></div>
<pre class="instructions" id="p-{gem["n"]:02d}">{esc(texts[gem["slug"]])}</pre>
</article>''')
    body = f'''<div class="wrap sheet">
<p class="kicker no-print"><span class="pill">Printable view</span></p>
<h1>All 33 desks</h1>
<p class="lede">This page is the field guide. Use the button to print, or copy any instruction block.</p>
<p class="no-print"><button class="solid" type="button" data-print="true">Print this guide</button></p>
{"".join(blocks)}
</div>'''
    return shell("Printable view · Gemini Gems Gallery", "Printable instructions for all 33 Gemini Gems.", 0, body)


def render_md(gem: dict, text: str) -> str:
    nodes = []
    for i, (title, _detail) in enumerate(gem["steps"]):
        nodes.append(f'  s{i}["{title.replace(chr(34), "")}"]')
    links = " --> ".join(f"s{i}" for i in range(len(gem["steps"])))
    mermaid = "flowchart LR\n" + "\n".join(nodes) + f"\n  {links}\n"
    steps = "\n".join(f"{i}. **{title}.** {detail}" for i, (title, detail) in enumerate(gem["steps"], 1))
    kid = "\n\n**Child data:** first name and age band only.\n" if gem.get("kid") else "\n"
    return f'''# {gem["icon"]} {gem["n"]:02d} · {gem["name"]}

**Category:** {gem["cat"]} · **Mark:** {gem["mark"]}

{gem["summary"]}

[Animated desk page](../site/gems/{fname(gem)}.html) · [Gallery](../site/index.html) · [Printable view](../site/print.html)

## How it works

{steps}

```mermaid
{mermaid}```

## Example input

```example
{gem["example_in"]}
```

## Example output

```example
{gem["example_out"]}
```

## Remix

{gem["remix"]}
{kid}
## Instructions

Copy the block below into [Gemini](https://gemini.google.com), then open **Gems > New Gem**.

```text
{text}```

## Use it

1. Open [Gemini](https://gemini.google.com) and go to **Gems > New Gem**.
2. Name it **{gem["name"]}** and paste the instructions.
3. Save, then try a prompt from the animated page.
4. To call it from code, read the [examples guide](../examples/README.md).

Downloads: [PDF guide](../downloads/Gemini-Gems-Guide.pdf) · [ZIP pack](../downloads/gemini-gems-pack.zip) · [Overview deck](../downloads/Gemini-Gems-Overview.pptx)

[Back to the gallery](../README.md)
'''


def render_readme(gems: list[dict]) -> str:
    groups = []
    for cat in CATS:
        rows = ["| | Desk | Page | What it does |", "|---|---|---|---|"]
        for gem in gems:
            if gem["cat"] != cat:
                continue
            rows.append(
                f'| {gem["icon"]} | [{gem["name"]}](gems/{fname(gem)}.md) | [Open](site/gems/{fname(gem)}.html) | {gem["summary"]} |'
            )
        anchor = cat.lower().replace(" & ", "-").replace(" ", "-")
        groups.append(f'<a id="{anchor}"></a>\n\n### {cat}\n\n' + "\n".join(rows))
    return f'''<div align="center">

# Gemini Gems Gallery

33 copy-ready Gemini Gems for work, engineering, learning, and home.

[![Gems](https://img.shields.io/badge/Gems-33-6366F1?style=for-the-badge)](#the-gems)
[![Download PDF](https://img.shields.io/badge/Download-PDF-0F9B8E?style=for-the-badge)](downloads/Gemini-Gems-Guide.pdf)
[![Download ZIP](https://img.shields.io/badge/Download-ZIP-22C55E?style=for-the-badge)](downloads/gemini-gems-pack.zip)
[![Download PPTX](https://img.shields.io/badge/Download-PPTX-6366F1?style=for-the-badge)](downloads/Gemini-Gems-Overview.pptx)
[![Gallery](https://img.shields.io/badge/Gallery-offline-F59E0B?style=for-the-badge)](site/index.html)
[![License](https://img.shields.io/badge/License-CC_BY_4.0_%2B_MIT-0B1220?style=for-the-badge)](LICENSE)

</div>

<a id="the-gems"></a>

## The desks

Each file in [gems/](gems/01-sales-call-prep-desk.md) is one Gem: a role, the inputs it needs, the steps it follows, the output shape, and the rules it will not break. The [offline gallery](site/index.html) adds search, category filters, copy buttons, a light theme, a dark theme, and a [printable view](site/print.html).

Three rules hold for the whole set:

- **Draft only.** A Gem writes text. You send, book, pay, post, merge, and file.
- **No invented facts.** If a number, date, or quote is not in your material, the Gem marks it UNKNOWN.
- **Narrow child data.** Desks about a child ask for a first name and an age band, and they do not give final homework answers.

{chr(10).join(groups)}

<a id="install"></a>

## Install a Gem

1. Open a file in [gems/](gems/02-meeting-recap-desk.md) and copy the text fence.
2. Go to [gemini.google.com](https://gemini.google.com), then **Gems > New Gem**.
3. Name it, paste the instructions, and save.
4. Start with the example on that desk's [gallery page](site/index.html).

<a id="build-an-app"></a>

## Build an app with a Gem

```mermaid
flowchart LR
  A[Pick a Gem] --> B[Copy instructions]
  B --> C[Set GEMINI_API_KEY]
  C --> D[Run the starter]
  D --> E[Review the draft]
```

The starters live in [examples/README.md](examples/README.md). They read `GEMINI_API_KEY` from the environment. Get a key in [Google AI Studio](https://aistudio.google.com/apikey).

```bash
pip install -r scripts/requirements.txt
export GEMINI_API_KEY="your-key"
python examples/gem_app.py gems/02-meeting-recap-desk.md "Recap these notes: ..."
```

Rebuild this gallery from the catalog with [scripts/build_gallery.py](scripts/build_gallery.py). Check links with [scripts/check_links.py](scripts/check_links.py). Fonts in the site and the PDF are [Inter](scripts/fonts/OFL.txt), under the SIL Open Font License.

## Downloads

| File | What it is |
|---|---|
| [Gemini-Gems-Guide.pdf](downloads/Gemini-Gems-Guide.pdf) | Designed guide to all 33 desks |
| [gemini-gems-pack.zip](downloads/gemini-gems-pack.zip) | Gallery, instructions, starters, guide, and deck |
| [Gemini-Gems-Overview.pptx](downloads/Gemini-Gems-Overview.pptx) | Overview deck with a link to every desk |

## Repository

```
gems/         one Markdown file per Gem
site/         offline gallery, one page per Gem
downloads/    PDF guide, ZIP pack, PowerPoint deck
examples/     Python and JavaScript starters
scripts/      generator, link checker, and fonts
```

To publish the gallery, serve the repository root and open [site/index.html](site/index.html).

## License

- Content (instructions, prompts, gallery text, the guide): [CC BY 4.0](https://creativecommons.org/licenses/by/4.0/)
- Code in [examples/](examples/README.md) and [site/](site/index.html): [MIT](LICENSE)

Gemini and Google AI Studio are trademarks of Google LLC. This project is not affiliated with Google.
'''


def render_examples() -> str:
    return '''# Build an app with a Gem

A Gem is a system prompt with a job and a set of rules. These starters load that prompt from a file in [gems/](../gems/02-meeting-recap-desk.md) and send one user message to the Gemini API.

```mermaid
flowchart LR
  A[Pick a Gem] --> B[Copy instructions]
  B --> C[Set GEMINI_API_KEY]
  C --> D[Run the starter]
  D --> E[Review the draft]
```

The same flow is on the [gallery page](../site/index.html#build-an-app).

## What the starter reads

Each Gem file has one fenced `text` block. That block is the whole instruction. The starter uses the first block and ignores the rest of the page. Do not add a second `text` fence.

## Set the key

Create a key in [Google AI Studio](https://aistudio.google.com/apikey). Put it in the environment. Do not write it into a source file.

```bash
export GEMINI_API_KEY="your-key"
```

Optional: `GEMINI_MODEL` overrides the default model, `gemini-2.5-flash`.

## Python

```bash
pip install google-genai
python examples/gem_app.py gems/02-meeting-recap-desk.md "Recap these notes: We will ship the export fix on Thursday. Priya owns the patch."
```

The script is [gem_app.py](gem_app.py). Run it from the repository root so the Gem path resolves.

## JavaScript

Node 18 or newer.

```bash
npm install @google/genai
node examples/gem_app.mjs gems/02-meeting-recap-desk.md "Recap these notes: We will ship the export fix on Thursday. Priya owns the patch."
```

The script is [gem_app.mjs](gem_app.mjs).

## What you get back

You get a draft in the terminal. The model is not connected to your mail, calendar, bank, or repository. If you want a different desk, change the file path. The meeting recap above is a safe sample because the notes are included in the prompt.

## More

- [Gallery](../site/index.html)
- [Printable view](../site/print.html)
- [PDF guide](../downloads/Gemini-Gems-Guide.pdf)
- [README](../README.md)
'''


class Bookmark(Flowable):
    def __init__(self, key: str, title: str):
        super().__init__()
        self.key = key
        self.title = title

    def wrap(self, aw, ah):
        return (0, 0)

    def draw(self):
        self.canv.bookmarkPage(self.key)
        self.canv.addOutlineEntry(self.title, self.key, level=0, closed=False)


def build_pdf(path: Path, gems: list[dict], texts: dict[str, str]) -> None:
    font_dir = ROOT / "scripts" / "fonts"
    pdfmetrics.registerFont(TTFont("Inter", str(font_dir / "Inter-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("InterBold", str(font_dir / "Inter-Bold.ttf")))
    navy = HexColor("#0B1220")
    teal = HexColor("#0F9B8E")
    indigo = HexColor("#6366F1")
    green = HexColor("#22C55E")
    amber = HexColor("#F59E0B")
    ink = HexColor("#142033")
    muted = HexColor("#3C4A63")

    def paint_cover(c, _doc):
        w, h = letter
        c.saveState()
        c.setFillColor(navy)
        c.rect(0, 0, w, h, fill=1, stroke=0)
        for i, col in enumerate((teal, indigo, green, amber)):
            c.setFillColor(col)
            c.rect(0, h - 18 - i * 12, w, 8, fill=1, stroke=0)
        c.setFillColor(HexColor("#E8EEF8"))
        c.setFont("InterBold", 34)
        c.drawString(54, h - 180, "Gemini Gems Gallery")
        c.setFont("Inter", 14)
        c.drawString(54, h - 210, "A designed guide to 33 copy-ready desks.")
        c.setFont("Inter", 11)
        y = h - 270
        for label, col in (("Workday", teal), ("Engineering", indigo), ("Learning", amber), ("Home", green)):
            c.setFillColor(col)
            c.roundRect(54, y - 4, 16, 16, 4, fill=1, stroke=0)
            c.setFillColor(HexColor("#E8EEF8"))
            c.drawString(80, y, label)
            y -= 28
        c.setFillColor(amber)
        c.roundRect(54, 78, 220, 32, 8, fill=1, stroke=0)
        c.setFillColor(navy)
        c.setFont("InterBold", 12)
        c.drawString(70, 90, "Open Gemini")
        c.linkURL("https://gemini.google.com", (54, 78, 274, 110), relative=0)
        c.setFillColor(HexColor("#8EA0C4"))
        c.setFont("Inter", 9)
        c.drawString(54, 48, "Drafts only. Not affiliated with Google.")
        c.restoreState()

    def paint_page(c, doc):
        w, h = letter
        c.saveState()
        c.setFillColor(navy)
        c.rect(0, h - 28, w, 28, fill=1, stroke=0)
        for i, col in enumerate((teal, indigo, green, amber)):
            c.setFillColor(col)
            c.rect(36 + i * 78, h - 10, 70, 5, fill=1, stroke=0)
        c.setFillColor(HexColor("#E8EEF8"))
        c.setFont("Inter", 9)
        c.drawString(360, h - 20, "Gemini Gems Gallery")
        c.setFillColor(muted)
        c.setFont("Inter", 8)
        c.drawString(36, 28, "Drafts only. Not financial, legal, or medical advice.")
        c.drawRightString(w - 36, 28, str(doc.page))
        c.restoreState()

    styles = {
        "h": ParagraphStyle("h", fontName="InterBold", fontSize=18, leading=22, textColor=ink, spaceAfter=6),
        "h2": ParagraphStyle("h2", fontName="InterBold", fontSize=12, leading=16, textColor=ink, spaceBefore=8, spaceAfter=3),
        "b": ParagraphStyle("b", fontName="Inter", fontSize=9.5, leading=13, textColor=ink, spaceAfter=4),
        "s": ParagraphStyle("s", fontName="Inter", fontSize=8.5, leading=11.4, textColor=ink),
        "link": ParagraphStyle("link", fontName="Inter", fontSize=10, leading=15, textColor=teal),
    }
    doc = BaseDocTemplate(
        str(path),
        pagesize=letter,
        title="Gemini Gems Gallery",
        author="Gemini Gems Gallery",
    )
    cover_frame = Frame(0, 0, 20, 20, id="cover", showBoundary=0)
    body = Frame(40, 46, letter[0] - 80, letter[1] - 86, id="body", showBoundary=0)
    doc.addPageTemplates([
        PageTemplate(id="cover", frames=[cover_frame], onPage=paint_cover),
        PageTemplate(id="body", frames=[body], onPage=paint_page),
    ])
    story = [NextPageTemplate("body"), PageBreak(), Bookmark("contents", "Contents")]
    story.append(Paragraph("Contents", styles["h"]))
    story.append(Paragraph("Each desk name jumps to its section in this file. The working instructions are included in full.", styles["b"]))
    for gem in gems:
        story.append(Paragraph(f'<a href="#g{gem["n"]:02d}" color="{CATS[gem["cat"]]}">{gem["n"]:02d}  {esc(gem["name"])}</a>  ·  {esc(gem["cat"])}', styles["link"]))
    story.append(Paragraph('<a href="https://gemini.google.com" color="#0F9B8E">Open Gemini</a>', styles["link"]))
    story.append(Paragraph('<a href="https://aistudio.google.com/apikey" color="#6366F1">Get an API key in Google AI Studio</a>', styles["link"]))
    story.append(Paragraph('<a href="https://creativecommons.org/licenses/by/4.0/" color="#0B1220">Content license: CC BY 4.0</a>', styles["b"]))
    for gem in gems:
        story.append(PageBreak())
        story.append(Bookmark(f'g{gem["n"]:02d}', f'{gem["n"]:02d} {gem["name"]}'))
        story.append(Paragraph(f'{gem["n"]:02d} · {esc(gem["name"])}', styles["h"]))
        story.append(Paragraph(f'{esc(gem["cat"])} · {esc(gem["mark"])}', styles["b"]))
        story.append(Paragraph(esc(gem["purpose"]), styles["b"]))
        story.append(Paragraph("How it works", styles["h2"]))
        for i, (title, detail) in enumerate(gem["steps"], 1):
            story.append(Paragraph(f"{i}. <b>{esc(title)}.</b> {esc(detail)}", styles["b"]))
        story.append(Paragraph("Instructions", styles["h2"]))
        block = esc(texts[gem["slug"]]).replace("\n", "<br/>")
        story.append(Paragraph(block, styles["s"]))
    doc.build(story)


def build_pptx(path: Path, gems: list[dict]) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Gemini Gems Gallery"
    prs.core_properties.author = "Gemini Gems Gallery"
    prs.core_properties.subject = "33 copy-ready Gemini Gems"
    blank = prs.slide_layouts[6]
    navy = RGBColor(0x0B, 0x12, 0x20)
    ink = RGBColor(0xE8, 0xEE, 0xF8)
    muted = RGBColor(0xB4, 0xC0, 0xD8)
    teal = RGBColor(0x0F, 0x9B, 0x8E)
    indigo = RGBColor(0x63, 0x66, 0xF1)
    green = RGBColor(0x22, 0xC5, 0x5E)
    amber = RGBColor(0xF5, 0x9E, 0x0B)

    def rgb(hex_color: str) -> RGBColor:
        value = hex_color.lstrip("#")
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))

    def slide():
        s = prs.slides.add_slide(blank)
        shape = s.shapes.add_shape(1, Emu(0), Emu(0), prs.slide_width, prs.slide_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = navy
        shape.line.fill.background()
        for i, col in enumerate((teal, indigo, green, amber)):
            bar = s.shapes.add_shape(1, Inches(0.4 + i * 0.55), Inches(0.28), Inches(0.42), Inches(0.12))
            bar.fill.solid()
            bar.fill.fore_color.rgb = col
            bar.line.fill.background()
        return s

    def text(s, left, top, width, height, value, size=20, color=ink, bold=False):
        box = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        run = frame.paragraphs[0].add_run()
        run.text = value
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return run

    s = slide()
    text(s, 0.5, 1.5, 12, 1.2, "Gemini Gems Gallery", 48, ink, True)
    text(s, 0.5, 2.8, 10, 1, "33 copy-ready desks for work, engineering, learning, and home.", 22, muted)
    text(s, 0.5, 4.2, 6, 0.4, "Draft only. No invented facts. Narrow child data.", 18, amber, True)
    run = text(s, 0.5, 5.6, 5, 0.4, "Open the offline gallery", 20, teal, True)
    run.hyperlink.address = "../site/index.html"
    run = text(s, 6.2, 5.6, 5, 0.4, "Open Gemini", 20, amber, True)
    run.hyperlink.address = "https://gemini.google.com"

    s = slide()
    text(s, 0.5, 0.7, 12, 0.6, "Three rules", 36, ink, True)
    rules = [
        (teal, "Draft only", "The desk writes. You send, book, pay, post, merge, and file."),
        (indigo, "No invented facts", "Missing names, numbers, dates, and quotes stay UNKNOWN."),
        (green, "Narrow child data", "A child's desk asks for a first name and an age band."),
    ]
    for i, (col, title, body) in enumerate(rules):
        card = s.shapes.add_shape(1, Inches(0.5 + i * 4.2), Inches(2.0), Inches(3.9), Inches(3.4))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0x14, 0x1C, 0x30)
        card.line.color.rgb = col
        text(s, 0.7 + i * 4.2, 2.3, 3.5, 0.8, title, 24, col, True)
        text(s, 0.7 + i * 4.2, 3.3, 3.5, 1.6, body, 16, ink)

    s = slide()
    text(s, 0.5, 0.7, 12, 0.6, "Four rooms", 36, ink, True)
    for i, (cat, color) in enumerate(CATS.items()):
        n = sum(1 for g in gems if g["cat"] == cat)
        card = s.shapes.add_shape(1, Inches(0.5), Inches(1.6 + i * 1.3), Inches(12.2), Inches(1.1))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0x14, 0x1C, 0x30)
        card.line.fill.background()
        text(s, 0.8, 1.75 + i * 1.3, 8, 0.4, cat, 24, rgb(color), True)
        text(s, 9.5, 1.8 + i * 1.3, 2.5, 0.4, f"{n} desks", 20, ink, True)

    for cat, color in CATS.items():
        s = slide()
        text(s, 0.5, 0.55, 12, 0.5, cat, 32, rgb(color), True)
        col_x = 0.5
        y = 1.4
        for gem in [g for g in gems if g["cat"] == cat]:
            run = text(s, col_x, y, 6, 0.36, f'{gem["n"]:02d}  {gem["name"]}', 18, ink, True)
            run.hyperlink.address = f'../gems/{fname(gem)}.md'
            y += 0.48
            if y > 6.6:
                col_x = 6.8
                y = 1.4

    s = slide()
    text(s, 0.5, 0.7, 12, 0.6, "Build an app with a Gem", 36, ink, True)
    steps = [
        "1  Pick a Gem from the gallery.",
        "2  Copy the instruction block.",
        "3  Set GEMINI_API_KEY in the environment.",
        "4  Run the Python or JavaScript starter.",
        "5  Review the draft. The app does not send it.",
    ]
    for i, line in enumerate(steps):
        text(s, 0.6, 1.7 + i * 0.8, 11, 0.5, line, 22, ink)
    run = text(s, 0.6, 6.2, 6, 0.4, "Open the examples guide", 20, teal, True)
    run.hyperlink.address = "../examples/README.md"
    run = text(s, 6.5, 6.2, 5, 0.4, "Get an API key", 20, amber, True)
    run.hyperlink.address = "https://aistudio.google.com/apikey"

    s = slide()
    text(s, 0.5, 0.7, 12, 0.6, "Install a Gem", 36, ink, True)
    for i, line in enumerate([
        "Open the desk file in gems/.",
        "Copy the instruction block.",
        "In Gemini, open Gems, then New Gem.",
        "Paste, name the Gem, and save.",
    ]):
        text(s, 0.6, 1.8 + i * 0.9, 11, 0.5, f"{i + 1}   {line}", 24, ink)
    run = text(s, 0.6, 6.2, 5, 0.4, "Open Gemini", 20, teal, True)
    run.hyperlink.address = "https://gemini.google.com"

    s = slide()
    text(s, 0.5, 0.7, 12, 0.6, "Downloads", 36, ink, True)
    run = text(s, 0.6, 2.0, 10, 0.5, "PDF guide  ·  Gemini-Gems-Guide.pdf", 24, teal, True)
    run.hyperlink.address = "Gemini-Gems-Guide.pdf"
    run = text(s, 0.6, 3.0, 10, 0.5, "ZIP pack  ·  gemini-gems-pack.zip", 24, green, True)
    run.hyperlink.address = "gemini-gems-pack.zip"
    run = text(s, 0.6, 4.0, 10, 0.5, "This deck stays beside those files.", 20, muted)
    run = text(s, 0.6, 5.2, 10, 0.5, "Content license: CC BY 4.0", 20, amber, True)
    run.hyperlink.address = "https://creativecommons.org/licenses/by/4.0/"

    s = slide()
    text(s, 0.5, 2.2, 12, 1, "Copy a desk. Keep the rules.", 36, ink, True)
    text(s, 0.5, 3.4, 12, 1, "The Gem drafts. You decide what leaves the page.", 20, muted)
    run = text(s, 0.5, 5.2, 6, 0.4, "Back to the gallery", 20, teal, True)
    run.hyperlink.address = "../site/index.html"
    prs.save(str(path))


def build_zip(path: Path) -> None:
    allowed = {"README.md", "LICENSE", "gems", "site", "examples", "scripts", "downloads"}
    skip = {"downloads/gemini-gems-pack.zip", "downloads/link-check.txt"}
    files = []
    for item in ROOT.rglob("*"):
        if not item.is_file():
            continue
        rel = item.relative_to(ROOT).as_posix()
        top = rel.split("/", 1)[0]
        if top not in allowed or rel in skip or "__pycache__" in rel or rel.endswith(".pyc"):
            continue
        files.append(rel)
    note = io.BytesIO()
    with zipfile.ZipFile(note, "w", compression=zipfile.ZIP_DEFLATED) as inner:
        inner.writestr(
            "README.txt",
            "These folders are the Gemini Gems pack.\n"
            "You already have the desks, the gallery, the PDF guide, and the deck.\n"
            "This small archive exists so the pack's own download link still opens after you unpack it.\n",
        )
    if path.exists():
        path.unlink()
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for rel in sorted(files):
            archive.write(ROOT / rel, rel)
        archive.writestr("downloads/gemini-gems-pack.zip", note.getvalue())


def main() -> None:
    gems = sorted(GEMS, key=lambda g: g["n"])
    texts = {g["slug"]: render_instructions(g) for g in gems}
    validate(gems, texts)
    gem_dir = ROOT / "gems"
    page_dir = ROOT / "site" / "gems"
    page_dir.mkdir(parents=True, exist_ok=True)
    (ROOT / "downloads").mkdir(exist_ok=True)
    expected_md = {f"{fname(g)}.md" for g in gems}
    expected_html = {f"{fname(g)}.html" for g in gems}
    for old in gem_dir.glob("*.md"):
        if old.name not in expected_md:
            old.unlink()
    for old in page_dir.glob("*.html"):
        if old.name not in expected_html:
            old.unlink()
    for gem in gems:
        (gem_dir / f"{fname(gem)}.md").write_text(render_md(gem, texts[gem["slug"]]), encoding="utf-8")
        (page_dir / f"{fname(gem)}.html").write_text(render_gem_page(gem, gems, texts[gem["slug"]]), encoding="utf-8")
    (ROOT / "site" / "index.html").write_text(render_index(gems, texts), encoding="utf-8")
    (ROOT / "site" / "print.html").write_text(render_print(gems, texts), encoding="utf-8")
    (ROOT / "README.md").write_text(render_readme(gems), encoding="utf-8")
    (ROOT / "examples" / "README.md").write_text(render_examples(), encoding="utf-8")
    pdf = ROOT / "downloads" / "Gemini-Gems-Guide.pdf"
    pptx = ROOT / "downloads" / "Gemini-Gems-Overview.pptx"
    pack = ROOT / "downloads" / "gemini-gems-pack.zip"
    build_pdf(pdf, gems, texts)
    build_pptx(pptx, gems)
    build_zip(pack)
    print(f"gems {len(gems)} pdf {pdf.stat().st_size} pptx {pptx.stat().st_size} zip {pack.stat().st_size}")


if __name__ == "__main__":
    main()
